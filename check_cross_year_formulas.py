"""Read-only independent evaluation of the generated ledger's SUM and arithmetic formulas."""
import sys,json,re,functools
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from openpyxl.utils.cell import range_boundaries,get_column_letter
from pathlib import Path
p=Path('跨年度測試'); w=openpyxl.load_workbook(p/'測試示意_115學年度產出.xlsx')
expected=json.loads((p/'expected.json').read_text(encoding='utf-8'))
@functools.lru_cache(None)
def value(sheet,ref):
    v=w[sheet][ref].value
    if v is None:return 0
    if isinstance(v,(float,int)):return v
    if not v.startswith('='): raise ValueError((sheet,ref,v))
    expr=v[1:]
    def sum_range(m):
        left,top,right,bottom=range_boundaries(m[1])
        return str(sum(value(sheet,f'{get_column_letter(c)}{r}') for r in range(top,bottom+1) for c in range(left,right+1)))
    expr=re.sub(r'SUM\(([A-Z]+\d+:[A-Z]+\d+)\)',sum_range,expr,flags=re.I)
    expr=re.sub(r"(?:'([^']+)'|([\w]+))!([A-Z]+\d+)",lambda m:str(value(m[1] or m[2],m[3])),expr)
    expr=re.sub(r'[A-Z]+\d+',lambda m:str(value(sheet,m[0])),expr)
    if not re.fullmatch(r'[\d\s.+\-*/()]+',expr):raise ValueError(expr)
    return eval(expr,{'__builtins__':{}},{})
report=[]
for month,exp in expected['months'].items():
    name=f'{int(month):02}分類帳';s=w[name]
    cum=next(c.row for row in s for c in row if c.column==5 and c.value=='截至本月份累計數')
    subtotal=cum-1; blank=subtotal-1
    actual=value(name,f'P{cum}')
    assert actual==exp['balance'],(name,actual,exp['balance'])
    assert value(name,'F4')==exp['opening'],(name,'opening')
    assert value(name,f'F{subtotal}')==exp['income']+exp['other']-exp['refund']
    assert value(name,f'O{subtotal}')==exp['expense']
    assert value(name,f'P{blank}')==actual,(name,'last blank balance',value(name,f'P{blank}'),actual)
    report.append({'sheet':name,'opening':value(name,'F4'),'balance':actual,'balanceCell':f'P{cum}'})
print(json.dumps(report,ensure_ascii=False,indent=2))
(p/'公式驗證結果.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
